# ia_ui.py
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import requests
import threading
import datetime
import os
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import subprocess
import sys

class IAAsistenteUI:
    def __init__(self, app_controller):
        self.app_controller = app_controller
        self.db_crm = self.app_controller.db_crm

    def open_reformular_hechos_dialog(self, caso_actual=None):
        """Abrir diálogo para reformular hechos con IA"""
        caso_actual_id = caso_actual['id'] if caso_actual else None
        caso_actual_caratula = caso_actual.get('caratula', "General") if caso_actual else "General"

        dialog = tk.Toplevel(self.app_controller.root)
        dialog.title(f"Reformular Hechos con IA (Caso: {caso_actual_caratula[:30]})")
        dialog.transient(self.app_controller.root)
        dialog.grab_set()
        dialog.geometry("700x600") 
        dialog.resizable(True, True)

        main_frame = ttk.Frame(dialog, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)
        main_frame.columnconfigure(0, weight=1)
        main_frame.rowconfigure(1, weight=2)  # Para el input_text_frame
        main_frame.rowconfigure(3, weight=3)  # Para el output_text_frame

        ttk.Label(main_frame, text="Ingrese los hechos del cliente (o texto a reformular):").grid(row=0, column=0, sticky=tk.NW, pady=(0,2))
        
        input_text_frame = ttk.Frame(main_frame)
        input_text_frame.grid(row=1, column=0, sticky='nsew', pady=2)
        input_text_frame.columnconfigure(0, weight=1)
        input_text_frame.rowconfigure(0, weight=1)
        
        hechos_entrada_text = tk.Text(input_text_frame, wrap=tk.WORD, height=10)
        hechos_entrada_text.grid(row=0, column=0, sticky='nsew')
        
        hechos_entrada_scroll = ttk.Scrollbar(input_text_frame, command=hechos_entrada_text.yview)
        hechos_entrada_scroll.grid(row=0, column=1, sticky='ns')
        hechos_entrada_text['yscrollcommand'] = hechos_entrada_scroll.set

        ttk.Label(main_frame, text="Hechos Reformulados por IA:").grid(row=2, column=0, sticky=tk.NW, pady=(5,2))
        
        output_text_frame = ttk.Frame(main_frame)
        output_text_frame.grid(row=3, column=0, sticky='nsew', pady=2)
        output_text_frame.columnconfigure(0, weight=1)
        output_text_frame.rowconfigure(0, weight=1)
        
        resultado_ia_text = tk.Text(output_text_frame, wrap=tk.WORD, height=15, state=tk.DISABLED)
        resultado_ia_text.grid(row=0, column=0, sticky='nsew')
        
        resultado_ia_scroll = ttk.Scrollbar(output_text_frame, command=resultado_ia_text.yview)
        resultado_ia_scroll.grid(row=0, column=1, sticky='ns')
        resultado_ia_text['yscrollcommand'] = resultado_ia_scroll.set

        status_var = tk.StringVar(value="Listo para recibir hechos.")
        status_label = ttk.Label(main_frame, textvariable=status_var, relief=tk.SUNKEN, anchor=tk.W)
        status_label.grid(row=4, column=0, sticky=tk.EW, pady=(5,5))

        # Botones
        buttons_frame = ttk.Frame(main_frame)
        buttons_frame.grid(row=5, column=0, sticky=tk.EW, pady=(5,0))

        procesar_btn = ttk.Button(buttons_frame, text="Reformular con IA", 
                                 command=lambda: self._solicitar_reformulacion(
                                     hechos_entrada_text, resultado_ia_text, status_var, 
                                     copiar_btn, guardar_docx_btn, dialog))
        procesar_btn.pack(side=tk.LEFT, padx=(0, 10))

        copiar_btn = ttk.Button(buttons_frame, text="Copiar Resultado", state=tk.DISABLED,
                               command=lambda: self._copiar_resultado_ia(resultado_ia_text, status_var))
        copiar_btn.pack(side=tk.LEFT, padx=(0, 10))

        guardar_docx_btn = ttk.Button(buttons_frame, text="Guardar como DOCX", state=tk.DISABLED,
                                     command=lambda: self._guardar_resultado_como_docx(
                                         resultado_ia_text, caso_actual, dialog, status_var))
        guardar_docx_btn.pack(side=tk.LEFT, padx=(0, 10))

        ttk.Button(buttons_frame, text="Cerrar", command=dialog.destroy).pack(side=tk.RIGHT)

        # Focus en el primer campo
        hechos_entrada_text.focus_set()

        return dialog

    def _solicitar_reformulacion(self, entrada_text, resultado_text, status_var, copiar_btn, guardar_btn, dialog):
        """Solicitar reformulación a la IA"""
        texto_hechos = entrada_text.get("1.0", tk.END).strip()
        if not texto_hechos:
            messagebox.showwarning("Entrada Vacía", "Por favor, ingrese el texto de los hechos a reformular.", parent=dialog)
            return

        status_var.set("Procesando con Asistente IA local, por favor espere...")
        resultado_text.config(state=tk.NORMAL)
        resultado_text.delete("1.0", tk.END)
        resultado_text.config(state=tk.DISABLED)
        dialog.update_idletasks()

        def do_request_thread():
            try:
                mcp_url = "http://localhost:5000/api/reformular_hechos"
                payload = {"texto_hechos": texto_hechos}
                response = requests.post(mcp_url, json=payload, timeout=90)
                response.raise_for_status()
                resultado_json = response.json()
                
                self.app_controller.root.after(0, lambda: self._actualizar_ui_con_respuesta(
                    resultado_json, resultado_text, status_var, copiar_btn, guardar_btn))
                    
            except requests.exceptions.ConnectionError:
                error_msg = (f"Error de Conexión: No se pudo conectar con el servidor del Asistente IA local en {mcp_url}.\n\n"
                           f"Verifique que:\n1. 'mcp_server.py' esté ejecutándose.\n"
                           f"2. Ollama/LM Studio esté activo y sirviendo el modelo correcto.\n"
                           f"3. No haya un firewall bloqueando la conexión a localhost en ese puerto.")
                self.app_controller.root.after(0, lambda: self._actualizar_ui_con_error(
                    error_msg, resultado_text, status_var, copiar_btn, guardar_btn, True))
                    
            except requests.exceptions.Timeout:
                error_msg = (f"Timeout: La solicitud al Asistente IA local en {mcp_url} tardó demasiado en responder (90s).\n\n"
                           f"Verifique el modelo LLM y la carga de su sistema.")
                self.app_controller.root.after(0, lambda: self._actualizar_ui_con_error(
                    error_msg, resultado_text, status_var, copiar_btn, guardar_btn))
                    
            except requests.exceptions.HTTPError as http_err:
                error_msg = f"Error HTTP {http_err.response.status_code} del servidor MCP: {http_err.response.text}"
                self.app_controller.root.after(0, lambda: self._actualizar_ui_con_error(
                    error_msg, resultado_text, status_var, copiar_btn, guardar_btn))
                    
            except requests.exceptions.JSONDecodeError:
                error_msg = "Error: El servidor MCP no devolvió una respuesta JSON válida."
                self.app_controller.root.after(0, lambda: self._actualizar_ui_con_error(
                    error_msg, resultado_text, status_var, copiar_btn, guardar_btn))
                    
            except Exception as e:
                error_msg = f"Error inesperado durante la solicitud a la IA: {type(e).__name__}: {e}"
                import traceback
                traceback.print_exc()
                self.app_controller.root.after(0, lambda: self._actualizar_ui_con_error(
                    error_msg, resultado_text, status_var, copiar_btn, guardar_btn))

        threading.Thread(target=do_request_thread, daemon=True).start()

    def _actualizar_ui_con_respuesta(self, resultado_json, resultado_text, status_var, copiar_btn, guardar_btn):
        """Actualizar UI con respuesta exitosa de la IA"""
        resultado_text.config(state=tk.NORMAL)
        resultado_text.delete("1.0", tk.END)
        
        if resultado_json and "hechos_reformulados" in resultado_json:
            resultado_text.insert("1.0", resultado_json["hechos_reformulados"])
            status_var.set("Respuesta de IA recibida.")
            copiar_btn.config(state=tk.NORMAL)
            guardar_btn.config(state=tk.NORMAL)
        elif resultado_json and "error" in resultado_json:
            error_msg_ia = f"Error devuelto por el Asistente IA: {resultado_json['error']}"
            resultado_text.insert("1.0", error_msg_ia)
            status_var.set("Error en la IA.")
            copiar_btn.config(state=tk.DISABLED)
            guardar_btn.config(state=tk.DISABLED)
        else:
            resultado_text.insert("1.0", "Respuesta inesperada o vacía del servidor.")
            status_var.set("Error: Respuesta no reconocida.")
            copiar_btn.config(state=tk.DISABLED)
            guardar_btn.config(state=tk.DISABLED)
            
        resultado_text.config(state=tk.DISABLED)

    def _actualizar_ui_con_error(self, mensaje_error, resultado_text, status_var, copiar_btn, guardar_btn, es_error_conexion=False):
        """Actualizar UI con error de comunicación"""
        resultado_text.config(state=tk.NORMAL)
        resultado_text.delete("1.0", tk.END)
        resultado_text.insert("1.0", f"Error en la comunicación:\n{mensaje_error}")
        resultado_text.config(state=tk.DISABLED)
        status_var.set("Error de comunicación.")
        
        if not es_error_conexion:
            messagebox.showerror("Error de Comunicación con IA", mensaje_error)
            
        copiar_btn.config(state=tk.DISABLED)
        guardar_btn.config(state=tk.DISABLED)

    def _copiar_resultado_ia(self, resultado_text, status_var):
        """Copiar resultado de IA al portapapeles"""
        texto_a_copiar = resultado_text.get("1.0", tk.END).strip()
        if texto_a_copiar:
            self.app_controller.root.clipboard_clear()
            self.app_controller.root.clipboard_append(texto_a_copiar)
            status_var.set("¡Resultado copiado al portapapeles!")
        else:
            messagebox.showwarning("Nada que Copiar", "No hay resultado para copiar.")

    def _guardar_resultado_como_docx(self, resultado_text, caso_actual, dialog, status_var):
        """Guardar resultado como documento DOCX"""
        texto_a_guardar = resultado_text.get("1.0", tk.END).strip()
        if not texto_a_guardar:
            messagebox.showwarning("Nada que Guardar", "No hay resultado para guardar.", parent=dialog)
            return

        caso_actual_caratula_saneada = "Hechos_IA"
        if caso_actual and caso_actual.get('caratula'):
            nombre_base = re.sub(r'[^\w\s-]', '', caso_actual.get('caratula', 'Caso'))
            nombre_base = re.sub(r'\s+', '_', nombre_base).strip('_')
            caso_actual_caratula_saneada = f"Hechos_IA_{nombre_base[:30]}"

        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        suggested_filename = f"{caso_actual_caratula_saneada}_{timestamp}.docx"
        
        filepath = filedialog.asksaveasfilename(
            title="Guardar Hechos como DOCX", 
            initialfile=suggested_filename, 
            defaultextension=".docx", 
            filetypes=[("Documento Word", "*.docx")], 
            parent=dialog
        )
        
        if filepath:
            try:
                doc = Document()
                doc.add_paragraph(texto_a_guardar)
                doc.save(filepath)
                messagebox.showinfo("Documento Guardado", f"Documento guardado en:\n{filepath}", parent=dialog)
                status_var.set(f"Guardado como {os.path.basename(filepath)}")
                
                if messagebox.askyesno("Abrir Documento", "¿Desea abrir el documento ahora?", parent=dialog):
                    self._abrir_archivo(filepath)
                    
                # Guardar interacción como actividad si hay caso asociado
                if caso_actual:
                    try:
                        self._guardar_interaccion_ia_como_actividad(
                            caso_actual['id'], 
                            "Reformulación IA", 
                            "Reformulación de hechos con IA", 
                            f"Documento guardado: {os.path.basename(filepath)}"
                        )
                    except Exception as e:
                        print(f"Error al guardar actividad: {e}")
                        
            except ImportError:
                messagebox.showerror("Error Librería", "Falta 'python-docx'. Instálala con: pip install python-docx", parent=dialog)
            except Exception as e_docx:
                messagebox.showerror("Error al Guardar DOCX", f"No se pudo guardar:\n{e_docx}", parent=dialog)

    def _abrir_archivo(self, filepath):
        """Abrir archivo con la aplicación predeterminada del sistema"""
        try:
            if sys.platform == "win32":
                os.startfile(filepath)
            elif sys.platform == "darwin":
                subprocess.call(["open", filepath])
            else:
                subprocess.call(["xdg-open", filepath])
        except Exception as e:
            print(f"Error al abrir archivo: {e}")

    def _guardar_interaccion_ia_como_actividad(self, caso_id, tipo_consulta, consulta, respuesta_ia):
        """Guardar interacción con IA como actividad de seguimiento"""
        if not caso_id:
            print("Advertencia: No se proporcionó caso_id para guardar actividad de IA.")
            return
        try:
            # Formatear la descripción completa de la interacción
            descripcion_completa = f"Tipo de Consulta IA: {tipo_consulta}\n\n"
            if consulta: # Si hay una consulta específica (puede no haberla si es solo una reformulación)
                 descripcion_completa += f"Consulta/Texto Original:\n{consulta}\n\n"
            descripcion_completa += f"Respuesta/Resultado de IA:\n{respuesta_ia}"

            fecha_hora_actual = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            # Usar add_actividad_caso con todos sus parámetros requeridos
            self.db_crm.add_actividad_caso(
                caso_id=caso_id,
                fecha_hora=fecha_hora_actual,
                tipo_actividad="Asistencia IA", # Un tipo de actividad genérico para IA
                descripcion=descripcion_completa,
                # creado_por y referencia_documento pueden ser None si no aplican aquí
                creado_por="Sistema IA",
                referencia_documento=None
            )
            print(f"Actividad de IA guardada para caso ID {caso_id}.")
        except Exception as e:
            print(f"Error al guardar actividad de IA para caso ID {caso_id}: {e}")
            import traceback
            traceback.print_exc()


    def open_sugerencia_caso_dialog(self, caso_actual=None):
        """Abrir diálogo para sugerencias de próximo paso (funcionalidad futura)"""
        if not caso_actual:
            messagebox.showinfo("Sin Caso", "Seleccione un caso para obtener sugerencias.")
            return

        dialog = tk.Toplevel(self.app_controller.root)
        dialog.title(f"Sugerencias IA para Caso: {caso_actual.get('caratula', 'Sin carátula')[:30]}")
        dialog.transient(self.app_controller.root)
        dialog.grab_set()
        dialog.geometry("600x500")
        dialog.resizable(True, True)

        main_frame = ttk.Frame(dialog, padding="20")
        main_frame.pack(fill=tk.BOTH, expand=True)
        main_frame.columnconfigure(0, weight=1)
        main_frame.rowconfigure(1, weight=1)

        # Información del caso
        info_frame = ttk.LabelFrame(main_frame, text="Información del Caso", padding="10")
        info_frame.grid(row=0, column=0, sticky='ew', pady=(0, 10))
        info_frame.columnconfigure(1, weight=1)

        ttk.Label(info_frame, text="Caso:").grid(row=0, column=0, sticky=tk.W, padx=(0, 10))
        ttk.Label(info_frame, text=caso_actual.get('caratula', 'N/A'), wraplength=400).grid(row=0, column=1, sticky=tk.EW)

        ttk.Label(info_frame, text="Etapa:").grid(row=1, column=0, sticky=tk.W, padx=(0, 10))
        ttk.Label(info_frame, text=caso_actual.get('etapa_procesal', 'N/A')).grid(row=1, column=1, sticky=tk.EW)

        # Área de sugerencias
        sugerencias_frame = ttk.LabelFrame(main_frame, text="Sugerencias de IA", padding="10")
        sugerencias_frame.grid(row=1, column=0, sticky='nsew', pady=(0, 10))
        sugerencias_frame.columnconfigure(0, weight=1)
        sugerencias_frame.rowconfigure(0, weight=1)

        sugerencias_text = tk.Text(sugerencias_frame, wrap=tk.WORD, state=tk.DISABLED)
        sugerencias_text.grid(row=0, column=0, sticky='nsew')

        sugerencias_scroll = ttk.Scrollbar(sugerencias_frame, command=sugerencias_text.yview)
        sugerencias_scroll.grid(row=0, column=1, sticky='ns')
        sugerencias_text['yscrollcommand'] = sugerencias_scroll.set

        # Mensaje de funcionalidad futura
        sugerencias_text.config(state=tk.NORMAL)
        sugerencias_text.insert("1.0", "🚧 Funcionalidad en desarrollo 🚧\n\n"
                                       "Esta función utilizará IA para analizar el estado actual del caso y sugerir próximos pasos basados en:\n\n"
                                       "• Etapa procesal actual\n"
                                       "• Tareas pendientes\n"
                                       "• Audiencias programadas\n"
                                       "• Actividades recientes\n"
                                       "• Plazos procesales\n\n"
                                       "Próximamente disponible...")
        sugerencias_text.config(state=tk.DISABLED)

        # Botones
        buttons_frame = ttk.Frame(main_frame)
        buttons_frame.grid(row=2, column=0, sticky='ew')

        ttk.Button(buttons_frame, text="Cerrar", command=dialog.destroy).pack(side=tk.RIGHT)

    def verificar_servidor_ia_disponible(self):
        """Verificar si el servidor de IA está disponible"""
        try:
            response = requests.get("http://localhost:5000/health", timeout=5)
            return response.status_code == 200
        except:
            return False

    def mostrar_estado_servidor_ia(self):
        """Mostrar estado del servidor de IA"""
        if self.verificar_servidor_ia_disponible():
            messagebox.showinfo("Estado del Servidor IA", 
                              "✅ Servidor de IA disponible\n\n"
                              "El asistente de IA está funcionando correctamente.")
        else:
            messagebox.showwarning("Estado del Servidor IA",
                                 "❌ Servidor de IA no disponible\n\n"
                                 "Verifique que:\n"
                                 "1. mcp_server.py esté ejecutándose\n"
                                 "2. Ollama/LM Studio esté activo\n"
                                 "3. El modelo correcto esté cargado")

class IAMenu:
    """Clase para manejar el menú de IA en la aplicación principal"""
    
    def __init__(self, menubar, app_controller):
        self.app_controller = app_controller
        self.ia_asistente = IAAsistenteUI(app_controller)
        
        # Crear menú de IA
        ia_menu = tk.Menu(menubar, tearoff=0)
        ia_menu.add_command(label="Reformular Hechos Cliente...", 
                           command=self._reformular_hechos_menu_callback)
        ia_menu.add_command(label="Sugerir Próximo Paso (Caso)...", 
                           command=self._sugerencia_caso_menu_callback)
        ia_menu.add_separator()
        ia_menu.add_command(label="Estado del Servidor IA", 
                           command=self.ia_asistente.mostrar_estado_servidor_ia)
        
        menubar.add_cascade(label="Asistente IA", menu=ia_menu)

    def _reformular_hechos_menu_callback(self):
        """Callback para el menú de reformular hechos"""
        caso_actual = None
        if hasattr(self.app_controller, 'selected_case') and self.app_controller.selected_case:
            caso_actual = self.app_controller.selected_case
        self.ia_asistente.open_reformular_hechos_dialog(caso_actual)

    def _sugerencia_caso_menu_callback(self):
        """Callback para el menú de sugerencias de caso"""
        caso_actual = None
        if hasattr(self.app_controller, 'selected_case') and self.app_controller.selected_case:
            caso_actual = self.app_controller.selected_case
        self.ia_asistente.open_sugerencia_caso_dialog(caso_actual)
